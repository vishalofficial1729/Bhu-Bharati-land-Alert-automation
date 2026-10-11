"""One monitoring pass: read Excel -> fetch each survey -> compare -> store -> alert -> report."""
import difflib
import json
import re
import smtplib
import sqlite3
import threading
import time
import urllib.request
from datetime import datetime
from email.message import EmailMessage
from pathlib import Path

from openpyxl import Workbook, load_workbook

from scraper import fetch_all

ROOT = Path(__file__).parent
DB = ROOT / "data" / "monitor.db"
REPORT = ROOT / "data" / "reports" / "land_status_report.xlsx"
FAILS_BEFORE_ALERT = 3
NAME_LABEL = "Pattadar / Authorised Person Name"
EXTENT_LABEL = "Total Extent (Ac. Gts)"
# Excel header -> internal key (matched by "header contains")
COLUMNS = {"district": "district", "mandal": "mandal", "village": "village", "survey": "survey",
           "owner": "owner", "acres": "acres", "guntas": "guntas", "alert on": "alert_on"}

# ---------- Telugu -> Latin (only good enough for fuzzy name matching) ----------
_V = dict(zip("అఆఇఈఉఊఋఎఏఐఒఓఔ", "a aa i ii u uu ru e ee ai o oo au".split()))
_C = dict(zip("కఖగఘఙచఛజఝఞటఠడఢణతథదధనపఫబభమయరఱలళవశషసహ",
              "k kh g gh n ch chh j jh n t th d dh n t th d dh n p ph b bh m y r r l l v sh sh s h".split()))
_S = dict(zip("ాిీుూృెేైొోౌ", "aa i ii u uu ru e ee ai o oo au".split()))


def telugu_to_latin(text: str) -> str:
    out = []
    for ch in text:
        if ch in _C:
            out.append(_C[ch] + "a")
        elif ch in _S and out and out[-1].endswith("a"):
            out[-1] = out[-1][:-1] + _S[ch]
        elif ch == "్" and out and out[-1].endswith("a"):
            out[-1] = out[-1][:-1]
        else:
            out.append(_V.get(ch) or {"ం": "m", "ః": "h", "ఁ": "n"}.get(ch, ch))
    return "".join(out)


def squash(name: str) -> str:
    # ponytail: crude phonetic squash; swap for a proper transliteration lib if matches get noisy.
    s = telugu_to_latin(name).lower()
    s = re.sub(r"[^a-z]", "", s).replace("h", "").replace("w", "v").replace("z", "j").replace("y", "i")
    return re.sub(r"(.)\1+", r"\1", s)


def name_score(expected: str, site_value: str) -> float:
    """Best fuzzy ratio vs any name on the site, ignoring word order."""
    def forms(n):
        return squash(n), "".join(sorted(squash(w) for w in n.split()))
    parts = [p for p in re.split(r"[/|]", site_value) if p.strip()]
    return max((difflib.SequenceMatcher(None, a, b).ratio()
                for p in parts for a, b in zip(forms(expected), forms(p))), default=0.0)


def site_extent_guntas(value: str):
    """'2.1525' (Ac.Gts) -> 2 acres 15.25 guntas -> 95.25 guntas."""
    m = re.match(r"\s*(\d+)(?:\.(\d+))?", value or "")
    if not m:
        return None
    dec = (m.group(2) or "").ljust(4, "0")
    return int(m.group(1)) * 40 + float(f"{dec[:2]}.{dec[2:]}")


# ---------- checks ----------
def derive(data: dict, row: dict) -> dict:
    """Computed fields (prefixed '_') that are always alerted on when they change."""
    out = {}
    vals = " ".join(data.values()).lower()
    flags = []
    tx = data.get("Transaction Status", "").strip()
    if tx and tx not in ("--", "-", "NA", "Nil", "No"):
        flags.append(f"Pending transaction ({tx})")
    if "prohibit" in vals:
        flags.append("Prohibited")
    status = data.get("Land Status", "")
    if status and ("digital" not in status.lower() or "not" in status.lower()):
        flags.append(f"Not digitally signed ({status})")
    if "nala" in vals or "non agri" in vals or "non-agri" in vals:
        flags.append("NALA / non-agricultural")
    out["_Flags"] = "; ".join(flags) or "None"

    if row.get("owner"):
        score = name_score(row["owner"], data.get(NAME_LABEL, ""))
        out["_Owner Check"] = "OK" if score >= 0.75 else f"MISMATCH: expected '{row['owner']}', site '{data.get(NAME_LABEL, '')}'"
    if row.get("acres") not in (None, "") or row.get("guntas") not in (None, ""):
        expected = float(row.get("acres") or 0) * 40 + float(row.get("guntas") or 0)
        site = site_extent_guntas(data.get(EXTENT_LABEL, ""))
        exp_txt = f"{int(expected // 40)} ac {expected % 40:g} gts"
        if site is None:
            out["_Extent Check"] = f"MISMATCH: expected {exp_txt}, site shows nothing"
        else:
            out["_Extent Check"] = "OK" if abs(site - expected) < 0.01 else \
                f"MISMATCH: expected {exp_txt}, site {int(site // 40)} ac {site % 40:g} gts ({data.get(EXTENT_LABEL)})"
    return out


def diff(old: dict, new: dict, watch: list) -> list:
    changes = []
    for k in sorted(set(old) | set(new)):
        if old.get(k) == new.get(k):
            continue
        if k.startswith("_") or not watch or any(w in k.lower() for w in watch):
            changes.append(f"{k.lstrip('_')}: '{old.get(k, '')}' -> '{new.get(k, '')}'")
    return changes


def first_run_issues(snap: dict) -> list:
    issues = [f"Flags: {snap['_Flags']}"] if snap.get("_Flags", "None") != "None" else []
    issues += [f"{k.lstrip('_')}: {v}" for k, v in snap.items() if k.endswith("Check") and v != "OK"]
    return issues


# ---------- storage ----------
def db():
    DB.parent.mkdir(exist_ok=True)
    con = sqlite3.connect(DB, timeout=30)
    con.executescript("""
        CREATE TABLE IF NOT EXISTS snapshots (key TEXT, ts TEXT, data TEXT);
        CREATE TABLE IF NOT EXISTS events (ts TEXT, key TEXT, kind TEXT, message TEXT);
        CREATE TABLE IF NOT EXISTS status (key TEXT PRIMARY KEY, last_checked TEXT, last_ok TEXT,
                                           fails INTEGER DEFAULT 0, last_error TEXT);
        CREATE TABLE IF NOT EXISTS runs (started TEXT, finished TEXT, total INTEGER, ok INTEGER,
                                         failed INTEGER, alerts INTEGER, seconds REAL);
        CREATE INDEX IF NOT EXISTS snapshots_key ON snapshots(key);
        CREATE INDEX IF NOT EXISTS events_key ON events(key);
        CREATE INDEX IF NOT EXISTS events_ts ON events(ts);""")
    if "seen" not in [c[1] for c in con.execute("PRAGMA table_info(events)")]:
        con.execute("ALTER TABLE events ADD COLUMN seen INTEGER DEFAULT 0")
    return con


def latest(con, key):
    r = con.execute("SELECT data FROM snapshots WHERE key=? ORDER BY rowid DESC LIMIT 1", (key,)).fetchone()
    return json.loads(r[0]) if r else None


def latest_all(con) -> dict:
    """Latest snapshot of every parcel in one query (rows are inserted in time order)."""
    return {k: json.loads(d) for k, d in con.execute(
        "SELECT key, data FROM snapshots WHERE rowid IN (SELECT MAX(rowid) FROM snapshots GROUP BY key)")}


# ---------- input ----------
def read_rows(path) -> list:
    wb = load_workbook(path, read_only=True, data_only=True)
    it = wb.worksheets[0].iter_rows(values_only=True)
    header = [str(h or "").strip().lower() for h in next(it)]
    idx = {}
    for i, h in enumerate(header):
        for needle, key in COLUMNS.items():
            if needle in h and key not in idx:
                idx[key] = i
    missing = [k for k in ("district", "mandal", "village", "survey") if k not in idx]
    if missing:
        raise ValueError(f"Excel is missing columns: {', '.join(missing)}")
    rows = []
    for n, r in enumerate(it, start=2):  # n = row number as shown in Excel
        row = {k: (str(r[i]).strip() if i < len(r) and r[i] is not None else "") for k, i in idx.items()}
        row["row"] = n
        if row["survey"]:
            row["survey"] = re.sub(r"\.0$", "", row["survey"])  # Excel turns 122 into 122.0
            row["key"] = f"{row['district']} / {row['mandal']} / {row['village']} / {row['survey']}"
            rows.append(row)
    wb.close()
    return rows


# ---------- output ----------
def send_email(cfg: dict, subject: str, body: str):
    s = cfg.get("email") or {}
    if not s.get("smtp_host") or not s.get("to"):
        raise RuntimeError("Email not configured (config.json -> email)")
    msg = EmailMessage()
    msg["Subject"], msg["From"] = subject, s.get("from") or s["username"]
    msg["To"] = ", ".join(s["to"]) if isinstance(s["to"], list) else s["to"]
    msg.set_content(body)
    port = int(s.get("smtp_port", 587))
    smtp = smtplib.SMTP_SSL(s["smtp_host"], port, timeout=30) if port == 465 else smtplib.SMTP(s["smtp_host"], port, timeout=30)
    with smtp:
        if port != 465:
            smtp.starttls()
        if s.get("username"):
            smtp.login(s["username"], s["password"])
        smtp.send_message(msg)


def current_table(con, rows):
    """One flat row per survey: fixed columns, then every site field seen."""
    all_snaps = latest_all(con)
    snaps = {r["key"]: all_snaps.get(r["key"], {}) for r in rows}
    stat = {k: (c, o, f, e) for k, c, o, f, e in con.execute("SELECT key,last_checked,last_ok,fails,last_error FROM status")}
    fields = []
    for s in snaps.values():
        fields += [k for k in s if k not in fields and not k.startswith("_")]
    head = ["District", "Mandal", "Village", "Survey No", "Pending Transaction", "Prohibited", "Digitally Signed",
            "NALA", "Owner Check", "Extent Check", "All Warnings", "Last Checked", "Last OK", "Error"] + fields
    table = [head]
    for r in rows:
        s = snaps[r["key"]]
        c, o, f, e = stat.get(r["key"], ("", "", 0, ""))
        flags = s.get("_Flags", "")
        yes = lambda needle: ("YES" if needle in flags else "No") if s else ""
        signed = ("No" if "Not digitally signed" in flags else "Yes") if s else ""
        table.append([r["district"], r["mandal"], r["village"], r["survey"], yes("Pending transaction"), yes("Prohibited"),
                      signed, yes("NALA"), s.get("_Owner Check", ""), s.get("_Extent Check", ""), flags,
                      c, o, e if f else ""] + [s.get(k, "") for k in fields])
    return table


def write_report(con, rows):
    REPORT.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Current"
    for r in current_table(con, rows):
        ws.append(r)
    ev = wb.create_sheet("History")
    ev.append(["Time", "Survey", "Type", "Details"])
    for r in con.execute("SELECT ts,key,kind,message FROM events ORDER BY ts DESC"):
        ev.append(list(r))
    for sh in (ws, ev):
        for col in sh.columns:
            sh.column_dimensions[col[0].column_letter].width = min(60, max(12, *(len(str(c.value or "")) for c in col)))
        sh.freeze_panes = "B2"
    wb.save(REPORT)


def push_google_sheet(cfg, con, rows, new_events):
    g = cfg.get("google_sheet") or {}
    if not g.get("web_app_url"):
        return
    payload = json.dumps({"token": g.get("token", ""), "current": current_table(con, rows),
                          "events": [[e["ts"], e["key"], e["kind"], e["message"]] for e in new_events]}).encode()
    req = urllib.request.Request(g["web_app_url"], payload, {"Content-Type": "application/json"})
    resp = urllib.request.urlopen(req, timeout=60).read().decode()
    if "ok" not in resp.lower():
        raise RuntimeError(f"Google Sheet push failed: {resp[:200]}")


# ---------- main pass ----------
PERMANENT = ("not found", "Multiple khatas")  # retrying these won't help


def run(cfg: dict, progress=lambda done, total, key: None, stop=None, only=None) -> dict:
    """Check every parcel in the sheet, or just the keys in `only` (re-check failed / one parcel)."""
    stop = stop or threading.Event()
    all_rows = list({r["key"]: r for r in read_rows(ROOT / cfg["excel_path"])}.values())  # drop duplicate rows
    rows = [r for r in all_rows if not only or r["key"] in only]
    by_loc = {(r["district"], r["mandal"], r["village"], r["survey"]): r for r in rows}
    con = db()
    now = lambda: datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    started, t0 = now(), time.time()
    prev_all = latest_all(con)
    new_events, done, ok, failed = [], 0, 0, 0

    def event(key, kind, message):
        e = {"ts": now(), "key": key, "kind": kind, "message": message}
        con.execute("INSERT INTO events(ts,key,kind,message,seen) VALUES (?,?,?,?,?)",
                    (*e.values(), 1 if kind == "BASELINE" else 0))
        new_events.append(e)

    def record_failure(row, msg):
        fails = con.execute("UPDATE status SET fails=fails+1, last_error=?, last_checked=? WHERE key=? RETURNING fails",
                            (msg, now(), row["key"])).fetchone()[0]
        if fails == FAILS_BEFORE_ALERT:
            event(row["key"], "ERROR", f"Failed {fails} checks in a row: {msg}")

    def record_success(row, data):
        key = row["key"]
        snap = {**data, **derive(data, row)}
        prev = prev_all.get(key)
        watch = [w.strip().lower() for w in row.get("alert_on", "").split(",") if w.strip()]
        if prev is None:
            for issue in first_run_issues(snap):
                event(key, "ISSUE", issue)
            event(key, "BASELINE", "First successful fetch stored")
        else:
            for c in diff(prev, snap, watch):
                event(key, "CHANGE", c)
        if snap != prev:
            con.execute("INSERT INTO snapshots VALUES (?,?,?)", (key, now(), json.dumps(snap, ensure_ascii=False)))
        con.execute("UPDATE status SET fails=0, last_error=NULL, last_checked=?, last_ok=? WHERE key=?", (now(), now(), key))

    con.executemany("INSERT OR IGNORE INTO status(key) VALUES (?)", [(r["key"],) for r in rows])
    con.commit()
    opts = dict(workers=int(cfg.get("parallel_browsers", 3)), twocaptcha_key=cfg.get("twocaptcha_api_key", ""),
                headless=not cfg.get("show_browser"), delay=float(cfg.get("delay_between_checks_sec", 1)), stop=stop)
    pending, retry = list(by_loc), {}
    for attempt in range(2):  # second pass retries the site's random "Something went wrong" errors
        for loc, res in fetch_all(pending, **opts):
            row = by_loc[loc]
            if isinstance(res, Exception) and attempt == 0 and not any(p in str(res) for p in PERMANENT):
                retry[loc] = res
                continue
            if isinstance(res, Exception):
                record_failure(row, str(res))
                failed += 1
            else:
                record_success(row, res)
                ok += 1
            con.commit()
            done += 1
            progress(done, len(rows), row["key"])
        pending = list(retry)
        retry = {}
        if not pending or stop.is_set():
            break
    for loc, row in ({} if stop.is_set() else by_loc).items():  # anything never returned (all browsers crashed)
        if con.execute("SELECT last_checked < ? OR last_checked IS NULL FROM status WHERE key=?", (started, row["key"])).fetchone()[0]:
            record_failure(row, "Not checked: the browser could not start")
            failed += 1
    con.commit()
    progress(len(rows), len(rows), "")

    write_report(con, all_rows)
    errors = []
    alertable = [e for e in new_events if e["kind"] != "BASELINE"]
    email_cfg = cfg.get("email") or {}
    if alertable and email_cfg.get("smtp_host") and email_cfg.get("to"):  # not set up yet = just skip
        body = "\n\n".join(f"[{e['kind']}] {e['key']}\n  {e['message']}" for e in alertable)
        try:
            send_email(cfg, f"Bhu Bharati alert: {len(alertable)} item(s) need attention", body +
                       f"\n\nFull report: {REPORT}")
        except Exception as ex:
            errors.append(f"Email failed: {ex}")
    try:
        push_google_sheet(cfg, con, all_rows, new_events)
    except Exception as ex:
        errors.append(f"Google Sheet failed: {ex}")
    con.execute("INSERT INTO runs VALUES (?,?,?,?,?,?,?)",
                (started, now(), len(rows), ok, failed, len(alertable), round(time.time() - t0, 1)))
    con.commit()
    con.close()
    return {"checked": done if stop.is_set() else len(rows), "total": len(rows), "ok": ok, "failed": failed,
            "alerts": len(alertable), "errors": errors, "seconds": round(time.time() - t0), "stopped": stop.is_set()}


if __name__ == "__main__":
    assert telugu_to_latin("పర్వతనేని అరుణ్") == "parvataneeni arun", telugu_to_latin("పర్వతనేని అరుణ్")
    assert name_score("Parvathaneni Arun", "పర్వతనేని అరుణ్ / పర్వతనేని అరుణ్") >= 0.75
    assert name_score("Induru Sudhakar Reddy", "ఇందూరు సుధాకర రెడ్డి / INDURU SUDHAKARA REDDY") >= 0.75
    assert name_score("Arun Parvathaneni", "పర్వతనేని అరుణ్") >= 0.75
    assert name_score("Test Owner One", "పర్వతనేని అరుణ్") < 0.75
    assert site_extent_guntas("0.0700") == 7 and site_extent_guntas("2.1525") == 95.25
    d = derive({"Transaction Status": "--", "Land Status": "Digital Signed", EXTENT_LABEL: "0.0700",
                NAME_LABEL: "పర్వతనేని అరుణ్"}, {"owner": "Parvathaneni Arun", "acres": "0", "guntas": "7"})
    assert d == {"_Flags": "None", "_Owner Check": "OK", "_Extent Check": "OK"}, d
    d = derive({"Transaction Status": "Mutation pending", "Land Status": "Not Digitally Signed"}, {})
    assert "Pending transaction" in d["_Flags"] and "Not digitally signed" in d["_Flags"], d
    assert diff({"A": "1", "B": "1"}, {"A": "2", "B": "2"}, ["a"]) == ["A: '1' -> '2'"]
    print("self-check ok")
